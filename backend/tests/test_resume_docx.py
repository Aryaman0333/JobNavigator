"""Editable Word export of a résumé (based on work by @mrAlishah, fork PR #2).

One plain layout for every template, the same content as the PDF in the PDF templates'
section order, and the footer only when the copy's PDF footer switch is on.
"""
import io
import uuid

import docx
import pytest
from docx.shared import Emu

from backend.accounts.context import use_account
from backend.api.routes_resumes import _build_docx, _docx_href, _resume_download_name, _xml_safe
from backend.models.db import Account, Resume
from backend.tests.r4_support import client, make_job  # noqa: F401

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CANDIDATE = "Dana Okonkwo"
DATA = {
    "header": {"name": CANDIDATE, "title": "Senior Product Manager",
               "contact_items": [{"text": "dana@example.com", "url": "mailto:dana@example.com"},
                                 {"text": "Berlin", "url": ""}]},
    "summary": "Product leader.",
    "experience": [{"title": "PM", "company": "Acme", "location": "Berlin", "date": "2020 - 2024",
                    "description": "Ran the platform.", "bullets": ["Shipped **Go** services", "Cut cost 20%"]}],
    "skills": {"Languages": "Go, Python"},
    "education": [{"school": "TU Berlin", "location": "Berlin", "degree": "BSc CS"}],
    "projects": [{"name": "Side", "description": "A thing", "bullets": ["Did it"]}],
    "publications": [],
}


def _resume(test_db, data=None, **over):
    over.setdefault("template", "inter")
    over.setdefault("is_base", True)
    over.setdefault("name", "Base CV")
    resume = Resume(id=uuid.uuid4(), page_format="letter", json_data=data or DATA, **over)
    test_db.add(resume)
    test_db.commit()
    return resume


def _copy(test_db, footer=None, **data_over):
    job = make_job(test_db, title="Staff Engineer", company="Acme", short_id=4242)
    data = {**DATA, **data_over}
    if footer is not None:
        data["footer_enabled"] = footer
    return _resume(test_db, data=data, is_base=False, job_id=job.id, name="Base CV → Acme")


def _fetch(client, resume, query=""):
    response = client.get(f"/api/resumes/{resume.id}/docx{query}")
    assert response.status_code == 200, response.text
    return response, docx.Document(io.BytesIO(response.content))


def _footer(document):
    return document.sections[0].footer.paragraphs[0].text


def _links(document):
    return [r.target_ref for r in document.part.rels.values() if r.reltype.endswith("/hyperlink")]


# ── the download ─────────────────────────────────────────────────────────────

def test_docx_download_headers_for_a_base_resume(client, test_db):
    response, _ = _fetch(client, _resume(test_db))
    assert response.headers["content-type"] == DOCX_TYPE
    assert response.content[:2] == b"PK"
    assert 'filename="DanaOkonkwo_BaseCV_Resume.docx"' in response.headers["content-disposition"]


def test_docx_filename_matches_the_pdf_one_for_a_copy(client, test_db):
    response, _ = _fetch(client, _copy(test_db))
    assert 'filename="DanaOkonkwo_BaseCV_Resume_4242.docx"' in response.headers["content-disposition"]


def test_download_name_tolerates_a_missing_header_and_drops_quotes(test_db):
    resume = Resume(id=uuid.uuid4(), name='My "CV"', is_base=True, json_data={"header": None})
    assert _resume_download_name(test_db, resume) == "Resume_MyCV_Resume"


def test_docx_page_size_follows_format(client, test_db):
    resume = _resume(test_db)
    _, letter = _fetch(client, resume)
    _, a4 = _fetch(client, resume, "?format=a4")
    assert letter.sections[0].page_width == Emu(7772400)
    assert a4.sections[0].page_width == Emu(7560310)


# ── content ──────────────────────────────────────────────────────────────────

def test_docx_contains_every_section_in_the_pdf_order(client, test_db):
    _, document = _fetch(client, _resume(test_db))
    text = [p.text for p in document.paragraphs]
    assert text[0] == CANDIDATE
    for heading in ("Experience", "Skills", "Education", "Projects"):
        assert heading in text
    assert "Publications" not in text   # an empty section prints no heading
    assert text.index("Experience") < text.index("Skills") < text.index("Education") < text.index("Projects")
    joined = "\n".join(text)
    for expected in ("Senior Product Manager", "Product leader.", "PM at Acme", "2020 - 2024", "Ran the platform.",
                     "Languages", "Go, Python", "TU Berlin", "BSc CS", "Side", "A thing"):
        assert expected in joined, expected


def test_docx_bullets_and_bold_markup(client, test_db):
    _, document = _fetch(client, _resume(test_db))
    bullets = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in bullets[:2]] == ["Shipped Go services", "Cut cost 20%"]
    assert [(r.text, bool(r.bold)) for r in bullets[0].runs] == [("Shipped ", False), ("Go", True), (" services", False)]


def test_docx_reads_legacy_dates_key_and_list_skills(client, test_db):
    legacy = {**DATA, "experience": [{"title": "PM", "company": "Acme", "dates": "1999 - 2001", "bullets": []}],
              "skills": [{"category": "Tools", "items": ["Jira", "SQL"]}]}
    _, document = _fetch(client, _resume(test_db, data=legacy))
    joined = "\n".join(p.text for p in document.paragraphs)
    assert "1999 - 2001" in joined
    assert "Tools: Jira, SQL" in joined


def test_docx_has_no_images_or_tables(client, test_db):
    _, document = _fetch(client, _resume(test_db))
    assert not document.inline_shapes and not document.tables


def test_docx_is_not_signed_by_the_library(client, test_db):
    _, document = _fetch(client, _resume(test_db))
    props = document.core_properties
    assert props.author == CANDIDATE
    assert "python-docx" not in f"{props.author}{props.last_modified_by}{props.comments}{props.title}"


def test_an_empty_resume_still_exports(client, test_db):
    _, document = _fetch(client, _resume(test_db, data={"header": {}}))
    assert document.paragraphs[0].text == ""


# ── links ────────────────────────────────────────────────────────────────────

def test_docx_contact_link_is_clickable(client, test_db):
    _, document = _fetch(client, _resume(test_db))
    assert _links(document) == ["mailto:dana@example.com"]


def test_phone_links_stay_tel_links(client, test_db):
    header = {"name": CANDIDATE, "contact_items": [
        {"text": "+1 555 0100", "url": "tel:+15550100"},
        {"text": "+1 (555) 0101", "url": "+1 (555) 0101"},
        {"text": "site", "url": "dana.dev"},
        {"text": "x", "url": "javascript:alert(1)"},
    ]}
    _, document = _fetch(client, _resume(test_db, data={**DATA, "header": header}))
    links = _links(document)
    assert "tel:+15550100" in links and "tel:+15550101" in links
    assert "https://dana.dev" in links
    assert not any(t.startswith("https://tel:") or t.startswith("javascript:") for t in links)
    assert "x" in document.paragraphs[1].text   # an unsafe scheme is printed, not linked


@pytest.mark.parametrize("url,href", [
    ("tel:+1555", "tel:+1555"),
    ("mailto:a@b.c", "mailto:a@b.c"),
    ("a@b.c", "mailto:a@b.c"),
    ("https://x.y/z", "https://x.y/z"),
    ("linkedin.com/in/dana", "https://linkedin.com/in/dana"),
    ("+44 20 7946 0958", "tel:+442079460958"),
    ("javascript:alert(1)", None),
    ("", None),
])
def test_docx_href(url, href):
    assert _docx_href(url) == href


# ── control characters (PDF imports) ─────────────────────────────────────────

def test_control_characters_never_break_the_export(client, test_db):
    data = {**DATA, "summary": "Led\x0bteams\x00 across\x1f regions\x0c", "header": {**DATA["header"], "title": "PM\x08"},
            "experience": [{"title": "PM", "company": "Acme", "bullets": ["Cut\x0bcost"]}]}
    _, document = _fetch(client, _resume(test_db, data=data))
    text = [p.text for p in document.paragraphs]
    assert "Led teams across regions " in text
    assert "PM" in text
    assert "Cut cost" in text


def test_xml_safe_keeps_tabs_and_newlines():
    assert _xml_safe({"a": ["x\ty\nz\r", 3, None]}) == {"a": ["x\ty\nz\r", 3, None]}
    assert _xml_safe("a\x0bb\x01c") == "a bc"


# ── footer: the PDF's opt-in switch ──────────────────────────────────────────

def test_copy_without_the_flag_has_no_footer(client, test_db):
    _, document = _fetch(client, _copy(test_db))
    assert _footer(document) == ""


def test_switched_off_footer_is_omitted(client, test_db):
    _, document = _fetch(client, _copy(test_db, footer=False))
    assert _footer(document) == ""


def test_switched_on_footer_is_printed(client, test_db):
    _, document = _fetch(client, _copy(test_db, footer=True))
    assert _footer(document) == f"{CANDIDATE} - Staff Engineer - Acme"


def test_base_resume_never_gets_a_footer(client, test_db):
    _, document = _fetch(client, _resume(test_db, data={**DATA, "footer_enabled": True}))
    assert _footer(document) == ""


def test_build_docx_without_footer_text_leaves_the_footer_empty():
    document = docx.Document(io.BytesIO(_build_docx(DATA, "letter")))
    assert _footer(document) == ""


# ── accounts ─────────────────────────────────────────────────────────────────

def test_docx_is_404_for_another_accounts_or_unknown_resume(client, test_db):
    b = Account(name="Partner", role="member")
    test_db.add(b)
    test_db.commit()
    with use_account(b.id):
        b_id = _resume(test_db).id
    test_db.expunge_all()
    for rid in (str(b_id), "not-a-uuid", str(uuid.uuid4())):
        r = client.get(f"/api/resumes/{rid}/docx")
        assert r.status_code == 404, (rid, r.status_code, r.text)
